"""
Master state schedule from the typed-in team schedules (assemble_schedules.py's inputs).

Every game listed on any team's page, once (a game both teams list is merged; "listed_by" says whose pages had it).
Scrimmages (before Oct 29) and placeholder opponents (TBA, tournament / playoff slots) are left out.
Rows are highlighted by how they relate to FOCUS_TEAM (Moorpark):
    light gold  - Moorpark's games
    pastel blue - two Moorpark opponents playing each other
    pastel red  - a Moorpark opponent against anyone else
Outputs, in "<season> Schedule Screenshots/":
    master_schedule.csv   (a "highlight" column, since CSV has no colors)
    master_schedule.xlsx  (colored rows, filterable)
    master_schedule.html  (colored rows, team filter)

Usage:
    python3 build_master_schedule.py [--season 2026-27] [--team Moorpark]
"""

import argparse
import csv
import html
import json
import re
from pathlib import Path

from assemble_schedules import canon, is_scrimmage, order
from ccc_config import DEFAULT_SEASON

PLACEHOLDER = re.compile(r"\b(TBA|TBD)\b|tourn|tourney|classic|tip ?off|playoff|round", re.I)
COLORS = {"moorpark": "FFE9A8", "opponents": "C6DDF5", "opponent": "F8C8C8", "": None}
LABELS = {"moorpark": "{t} game", "opponents": "{t} opponents vs each other", "opponent": "{t} opponent game", "": ""}
FIELDS = ["date", "time", "away", "home", "site", "conference", "notes", "listed_by", "box_status", "highlight"]


def iso(d, season):
    m, _, day = d.partition(" ")
    y = int(season[:4]) + (0 if m in ("Jul", "Aug", "Sep", "Oct", "Nov", "Dec") else 1)
    return f"{y}-{['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].index(m) + 1:02d}-{int(day):02d}"


def build(season, focus):
    base = Path(f"{season} Schedule Screenshots")
    own = {}
    for f in sorted(base.glob("*/*/schedule.json")):
        d = json.load(open(f))
        own[canon(d["team_name"]) or d["team_name"]] = d
    games = {}
    for team, d in own.items():
        for g in d["games"]:
            if is_scrimmage(g["date"]):
                continue
            opp = canon(g.get("opponent")) or g.get("opponent", "").strip()
            if not opp or (not canon(opp) and PLACEHOLDER.search(opp)):
                continue
            key = (g["date"], frozenset((team, opp)), g.get("time", ""))
            # Same pair, same day, listed by both sides: merge even if one side's time differs.
            match = next((k for k in games if k[0] == g["date"] and k[1] == key[1]), None)
            if match:
                r = games[match]
                r["listed_by"].append(team)
                if g.get("notes") and g["notes"] not in r["notes"]:
                    r["notes"].append(g["notes"])
                r["conference"] |= bool(g.get("is_conference"))
                continue
            ha = g.get("home_away")
            away, home = (opp, team) if ha == "home" else (team, opp)
            games[key] = {"date": g["date"], "time": g.get("time", ""), "away": away, "home": home,
                          "site": "neutral" if ha == "neutral" else "home", "conference": bool(g.get("is_conference")),
                          "notes": [g["notes"]] if g.get("notes") else [], "listed_by": [team]}
    opps = {canon(g["opponent"]) for g in own.get(focus, {}).get("games", []) if canon(g.get("opponent"))} - {focus}
    rows = []
    for r in sorted(games.values(), key=lambda r: (order(r), r["time"], r["away"])):
        t = {r["away"], r["home"]}
        hl = "moorpark" if focus in t else "opponents" if t <= opps else "opponent" if t & opps else ""
        rows.append({"date": iso(r["date"], season), "time": r["time"], "away": r["away"], "home": r["home"],
                     "site": r["site"], "conference": "Y" if r["conference"] else "", "notes": " | ".join(r["notes"]),
                     "listed_by": " + ".join(r["listed_by"]), "box_status": "SCHEDULED",
                     "highlight": LABELS[hl].format(t=focus), "_hl": hl})
    return base, rows, sorted(opps)


def write_csv(base, rows):
    with open(base / "master_schedule.csv", "w", newline="") as f:
        w = csv.DictWriter(f, FIELDS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(base, rows, focus):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Master schedule"
    ws.append(FIELDS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[k] for k in FIELDS])
        if COLORS[r["_hl"]]:
            fill = PatternFill("solid", fgColor=COLORS[r["_hl"]])
            for c in ws[ws.max_row]:
                c.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, wdt in zip("ABCDEFGHIJ", (11, 9, 20, 20, 8, 11, 40, 28, 12, 28)):
        ws.column_dimensions[col].width = wdt
    wb.save(base / "master_schedule.xlsx")


def write_html(base, rows, season, focus):
    teams = sorted({r["away"] for r in rows} | {r["home"] for r in rows})
    body = "\n".join(
        f'<tr class="{r["_hl"]}" data-t="{html.escape(r["away"])}|{html.escape(r["home"])}">'
        + "".join(f"<td>{html.escape(r[k])}</td>" for k in FIELDS[:-1]) + "</tr>" for r in rows)
    legend = "".join(f'<span class="{k}">{LABELS[k].format(t=focus)}</span>' for k in ("moorpark", "opponents", "opponent"))
    page = f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{season} Master Schedule</title><style>
body{{font:13px system-ui,sans-serif;margin:16px;background:#fff;color:#111}}
table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ddd;padding:4px 6px;text-align:left}}
th{{background:#222;color:#fff;position:sticky;top:0}}.wrap{{overflow-x:auto}}
.moorpark{{background:#{COLORS["moorpark"]}}}.opponents{{background:#{COLORS["opponents"]}}}.opponent{{background:#{COLORS["opponent"]}}}
.legend span{{display:inline-block;padding:3px 8px;margin:0 6px 8px 0;border:1px solid #ccc}}
</style></head><body><h2>{season} Master Schedule ({len(rows)} games)</h2>
<div class="legend">{legend}</div>
<label>Team <select id="f"><option value="">All teams</option>{"".join(f"<option>{html.escape(t)}</option>" for t in teams)}</select></label>
<div class="wrap"><table><thead><tr>{"".join(f"<th>{k}</th>" for k in FIELDS[:-1])}</tr></thead><tbody>
{body}
</tbody></table></div><script>
document.getElementById('f').onchange=e=>{{const v=e.target.value;document.querySelectorAll('tbody tr').forEach(tr=>{{
tr.style.display=!v||tr.dataset.t.split('|').includes(v)?'':'none'}})}};
</script></body></html>"""
    (base / "master_schedule.html").write_text(page)


def main():
    ap = argparse.ArgumentParser(description="Master state schedule (CSV, XLSX, HTML) from the typed-in schedules")
    ap.add_argument("--season", default=DEFAULT_SEASON)
    ap.add_argument("--team", default="Moorpark", help="team whose games / opponents are highlighted")
    a = ap.parse_args()
    base, rows, opps = build(a.season, a.team)
    write_csv(base, rows)
    write_xlsx(base, rows, a.team)
    write_html(base, rows, a.season, a.team)
    n = {k: sum(r["_hl"] == k for r in rows) for k in COLORS}
    print(f"{len(rows)} games ({a.team} {n['moorpark']}, opponents vs each other {n['opponents']}, "
          f"opponent games {n['opponent']}, other {n['']}); {len(opps)} {a.team} opponents -> {base}/master_schedule.csv/.xlsx/.html")


if __name__ == "__main__":
    main()
